#!/usr/bin/env python3
"""
scripts/excel_a_estate.py — mete datos propios (Excel o CSV) al pipeline.

El agente no lee Excel: lee una estate SQLite con el esquema de
docs/spec/estate_schema.sql. Este script es el puente. Convierte un
archivo tuyo en una estate valida, y a partir de ahi TODO lo demas
funciona igual que con las estates generadas: las reglas deterministas,
el CART, Gemma, el validador y el expediente.

    # 1. saca una plantilla con las columnas que se esperan
    python3 scripts/excel_a_estate.py --plantilla plantilla.xlsx

    # 2. llenala (o exporta tus datos con esos encabezados) y convierte
    python3 scripts/excel_a_estate.py --entrada mis_datos.xlsx \\
        --salida data/estates/mi_estate.db --puntuar

    # 3. corre el pipeline completo sobre ella
    python3 -m src.run --estate data/estates/mi_estate.db --out salida.json

FORMATOS QUE ACEPTA
  - .xlsx con una hoja por tabla (vendors, invoices, ledger, bank_txns,
    purchase_orders, contracts, employees, efos_list). Las hojas que
    falten quedan vacias.
  - una carpeta con CSVs nombrados igual (invoices.csv, vendors.csv...).
  - un solo CSV o una sola hoja de facturas:  --entrada facturas.csv --tabla invoices

Los encabezados se reconocen en espanol o en ingles y sin distinguir
acentos ni mayusculas: "RFC Emisor", "rfc_emisor" e "issuer_rfc" son la
misma columna. Lo que no se reconoce se reporta, no se adivina.

LO QUE ESTE SCRIPT *NO* HACE, A PROPOSITO
  No inventa evidencia. Si tu archivo no trae transferencias bancarias,
  la tabla bank_txns queda vacia y los detectores que dependen de ella
  (ciclos de transferencias, kickbacks, pagos no rastreables) no van a
  disparar — y el script te lo dice al final en vez de fabricar
  movimientos que nadie podria auditar. Un hallazgo que cita un
  record_id inexistente es justo lo que el validador esta puesto a
  rechazar.

  Lo unico que rellena es lo que se deduce sin inventar: los proveedores
  a partir de quien emitio cada factura, el IVA al 16% cuando das el
  subtotal, y el cruce contra el listado 69-B REAL del SAT
  (data/raw/Listado_completo_69-B.csv) para los RFC que aparezcan en tu
  archivo. Cada relleno se reporta.
"""

from __future__ import annotations

import argparse
import csv
import re
import sqlite3
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

SCHEMA_SQL = REPO / "docs" / "spec" / "estate_schema.sql"
SAT_CSV = REPO / "data" / "raw" / "Listado_completo_69-B.csv"
IVA_RATE = 0.16

# ---------------------------------------------------------------------------
# Esquema: columnas por tabla, y como se llama cada una en un Excel real.
# La primera columna de cada lista es la llave primaria.
# ---------------------------------------------------------------------------

TABLAS: dict[str, dict] = {
    "vendors": {
        "cols": ["rfc", "legal_name", "registered_date", "address", "bank_clabe",
                 "category", "contact_email"],
        "prefijo": None,
        "alias": {
            "rfc": ["rfc", "rfc proveedor", "rfc del proveedor", "vendor rfc"],
            "legal_name": ["razon social", "nombre", "nombre del contribuyente",
                           "proveedor", "legal name"],
            "registered_date": ["fecha de alta", "fecha alta", "fecha de registro",
                                "alta", "registered date"],
            "address": ["direccion", "domicilio", "address"],
            "bank_clabe": ["clabe", "cuenta", "cuenta bancaria", "bank clabe"],
            "category": ["categoria", "giro", "rubro", "category"],
            "contact_email": ["correo", "email", "e mail", "contacto", "contact email"],
        },
    },
    "invoices": {
        "cols": ["uuid", "issuer_rfc", "receiver_rfc", "issue_date", "subtotal",
                 "iva", "total", "concepto_text", "uso_cfdi", "forma_pago",
                 "metodo_pago", "status"],
        "prefijo": "INV",
        "alias": {
            "uuid": ["uuid", "folio fiscal", "folio", "id", "id factura",
                     "numero de factura", "invoice"],
            "issuer_rfc": ["rfc emisor", "emisor", "rfc del emisor", "issuer rfc",
                           "rfc proveedor"],
            "receiver_rfc": ["rfc receptor", "receptor", "rfc del receptor",
                             "receiver rfc", "rfc cliente"],
            "issue_date": ["fecha", "fecha de emision", "fecha emision",
                           "issue date", "date"],
            "subtotal": ["subtotal", "importe", "base"],
            "iva": ["iva", "impuesto", "impuestos", "vat"],
            "total": ["total", "monto", "monto total", "importe total"],
            "concepto_text": ["concepto", "descripcion", "concepto text",
                              "descripcion del servicio"],
            "uso_cfdi": ["uso cfdi", "uso del cfdi", "uso"],
            "forma_pago": ["forma de pago", "forma pago"],
            "metodo_pago": ["metodo de pago", "metodo pago"],
            "status": ["estatus", "estado", "status", "vigente"],
        },
    },
    "ledger": {
        "cols": ["entry_id", "date", "account_code", "account_name", "debit",
                 "credit", "description", "invoice_uuid", "cost_center", "approver"],
        "prefijo": None,
        "alias": {
            "entry_id": ["entry id", "poliza", "id", "numero de poliza", "asiento"],
            "date": ["fecha", "date"],
            "account_code": ["cuenta", "codigo de cuenta", "codigo contable",
                             "account code"],
            "account_name": ["nombre de cuenta", "cuenta contable", "account name"],
            "debit": ["debe", "cargo", "debit"],
            "credit": ["haber", "abono", "credit"],
            "description": ["descripcion", "concepto", "description"],
            "invoice_uuid": ["uuid", "factura", "folio fiscal", "invoice uuid"],
            "cost_center": ["centro de costo", "centro de costos", "cost center"],
            "approver": ["aprobador", "autorizo", "aprobo", "approver"],
        },
    },
    "bank_txns": {
        "cols": ["txn_id", "date", "from_clabe", "to_clabe", "amount", "reference",
                 "channel"],
        "prefijo": "BNK",
        "alias": {
            "txn_id": ["txn id", "id", "folio", "id movimiento", "movimiento"],
            "date": ["fecha", "fecha de la transferencia", "date"],
            "from_clabe": ["clabe origen", "cuenta origen", "origen", "de",
                           "from clabe", "clabe ordenante"],
            "to_clabe": ["clabe destino", "cuenta destino", "destino", "para",
                         "to clabe", "clabe beneficiario"],
            "amount": ["monto", "importe", "amount", "cantidad"],
            "reference": ["referencia", "concepto", "reference"],
            "channel": ["canal", "medio", "channel", "via"],
        },
    },
    "purchase_orders": {
        "cols": ["po_id", "vendor_rfc", "date", "amount", "requester", "approver",
                 "description"],
        "prefijo": "PO",
        "alias": {
            "po_id": ["po id", "orden de compra", "orden", "numero de orden", "folio", "id"],
            "vendor_rfc": ["rfc", "rfc proveedor", "proveedor", "vendor rfc"],
            "date": ["fecha", "date"],
            "amount": ["monto", "importe", "amount", "total"],
            "requester": ["solicitante", "solicito", "requester", "pidio"],
            "approver": ["aprobador", "autorizo", "aprobo", "approver"],
            "description": ["descripcion", "concepto", "description"],
        },
    },
    "contracts": {
        "cols": ["contract_id", "vendor_rfc", "start_date", "value", "scope_text"],
        "prefijo": "CTR",
        "alias": {
            "contract_id": ["contract id", "contrato", "numero de contrato", "id"],
            "vendor_rfc": ["rfc", "rfc proveedor", "proveedor", "vendor rfc"],
            "start_date": ["fecha de inicio", "inicio", "fecha", "start date"],
            "value": ["valor", "monto", "importe", "value"],
            "scope_text": ["alcance", "objeto", "descripcion", "scope text"],
        },
    },
    "employees": {
        "cols": ["emp_id", "name", "role", "bank_clabe", "hire_date"],
        "prefijo": "EMP",
        "alias": {
            "emp_id": ["emp id", "id", "numero de empleado", "empleado"],
            "name": ["nombre", "name"],
            "role": ["puesto", "rol", "cargo", "role"],
            "bank_clabe": ["clabe", "cuenta", "cuenta bancaria", "bank clabe"],
            "hire_date": ["fecha de ingreso", "ingreso", "fecha de contratacion",
                          "hire date"],
        },
    },
    "efos_list": {
        "cols": ["rfc", "legal_name", "status", "publication_date"],
        "prefijo": None,
        "alias": {
            "rfc": ["rfc"],
            "legal_name": ["razon social", "nombre", "nombre del contribuyente",
                           "legal name"],
            "status": ["situacion", "situacion del contribuyente", "estatus", "status"],
            "publication_date": ["fecha de publicacion", "publicacion",
                                 "publication date"],
        },
    },
}

NUMERICAS = {"subtotal", "iva", "total", "debit", "credit", "amount", "value"}
FECHAS = {"registered_date", "issue_date", "date", "start_date", "hire_date",
          "publication_date"}
RFC_RE = re.compile(r"^[A-ZÑ&]{3,4}\d{6}[A-Z0-9]{3}$")


# ---------------------------------------------------------------------------
# Normalizacion
# ---------------------------------------------------------------------------

def norm(s) -> str:
    """'RFC del Emisor ' -> 'rfc del emisor'. Sin acentos, sin puntuacion."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-zA-Z0-9]+", " ", s).strip().lower()
    return re.sub(r"\s+", " ", s)


def parse_fecha(v, avisos: list[str], ctx: str) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d",
                "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19] if " " in s else s, fmt).date().isoformat()
        except ValueError:
            continue
    avisos.append(f"{ctx}: fecha no reconocida '{s}', se guardo tal cual")
    return s


def parse_num(v, avisos: list[str], ctx: str) -> float | None:
    """Acepta 1234.56, '$1,234.56', '1 234,56' y celdas vacias."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else \
            s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", "") if len(s.split(",")[-1]) == 3 else s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        avisos.append(f"{ctx}: monto no numerico '{v}', se dejo vacio")
        return None


# ---------------------------------------------------------------------------
# Lectura de la entrada
# ---------------------------------------------------------------------------

def leer_xlsx(path: Path) -> dict[str, list[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit(
            "Para leer .xlsx hace falta openpyxl:\n"
            "    pip install openpyxl\n"
            "O exporta tus hojas a CSV y usa --entrada con una carpeta.")
    wb = load_workbook(path, data_only=True, read_only=True)
    hojas: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue
        encabezados = [h for h in filas[0]]
        datos = []
        for fila in filas[1:]:
            if all(c is None or str(c).strip() == "" for c in fila):
                continue
            datos.append({h: v for h, v in zip(encabezados, fila) if h is not None})
        hojas[ws.title] = datos
    wb.close()
    return hojas


def leer_csv(path: Path) -> list[dict]:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            with open(path, encoding=enc, newline="") as f:
                muestra = f.read(8192)
                f.seek(0)
                try:
                    dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
                except csv.Error:
                    dialecto = csv.excel
                return [r for r in csv.DictReader(f, dialect=dialecto)
                        if any((v or "").strip() for v in r.values())]
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"No pude decodificar {path} ni como UTF-8 ni como latin-1.")


def cargar_entrada(entrada: Path, tabla_forzada: str | None) -> dict[str, list[dict]]:
    if entrada.is_dir():
        hojas = {}
        for f in sorted(entrada.glob("*.csv")):
            hojas[f.stem] = leer_csv(f)
        if not hojas:
            raise SystemExit(f"No hay CSVs en {entrada}")
        return hojas
    if entrada.suffix.lower() in (".xlsx", ".xlsm"):
        hojas = leer_xlsx(entrada)
        if tabla_forzada:
            primera = next(iter(hojas.values()))
            return {tabla_forzada: primera}
        return hojas
    if entrada.suffix.lower() == ".csv":
        return {tabla_forzada or "invoices": leer_csv(entrada)}
    raise SystemExit(f"Extension no soportada: {entrada.suffix} (usa .xlsx, .csv o una carpeta)")


# ---------------------------------------------------------------------------
# Mapeo de columnas
# ---------------------------------------------------------------------------

def emparejar_hoja(nombre_hoja: str) -> str | None:
    n = norm(nombre_hoja)
    for tabla in TABLAS:
        if n == norm(tabla):
            return tabla
    equivalentes = {
        "facturas": "invoices", "cfdi": "invoices", "proveedores": "vendors",
        "poliza": "ledger", "polizas": "ledger", "contabilidad": "ledger",
        "mayor": "ledger", "bancos": "bank_txns", "transferencias": "bank_txns",
        "movimientos": "bank_txns", "estado de cuenta": "bank_txns",
        "ordenes de compra": "purchase_orders", "ordenes": "purchase_orders",
        "contratos": "contracts", "empleados": "employees", "nomina": "employees",
        "69 b": "efos_list", "lista 69 b": "efos_list", "efos": "efos_list",
    }
    return equivalentes.get(n)


def mapear_columnas(tabla: str, encabezados) -> tuple[dict, list[str]]:
    """Devuelve {columna_canonica: encabezado_original} y los no reconocidos."""
    spec = TABLAS[tabla]
    indice: dict[str, str] = {}
    for canon in spec["cols"]:
        indice[norm(canon)] = canon
        for a in spec["alias"].get(canon, []):
            indice.setdefault(norm(a), canon)

    mapa: dict[str, str] = {}
    sobrantes: list[str] = []
    for h in encabezados:
        canon = indice.get(norm(h))
        if canon and canon not in mapa:
            mapa[canon] = h
        elif canon is None:
            sobrantes.append(str(h))
    return mapa, sobrantes


def convertir_tabla(tabla: str, filas: list[dict], avisos: list[str]) -> list[dict]:
    if not filas:
        return []
    spec = TABLAS[tabla]
    mapa, sobrantes = mapear_columnas(tabla, filas[0].keys())
    if not mapa:
        avisos.append(f"{tabla}: ninguna columna reconocida; la hoja se ignoro")
        return []
    if sobrantes:
        avisos.append(f"{tabla}: columnas ignoradas (no estan en el esquema): "
                      + ", ".join(sobrantes[:8]))

    pk = spec["cols"][0]
    salida = []
    for i, fila in enumerate(filas, 1):
        ctx = f"{tabla} fila {i}"
        r: dict = {}
        for canon in spec["cols"]:
            v = fila.get(mapa[canon]) if canon in mapa else None
            if isinstance(v, str):
                v = v.strip() or None
            if canon in FECHAS:
                v = parse_fecha(v, avisos, ctx)
            elif canon in NUMERICAS:
                v = parse_num(v, avisos, ctx)
            elif canon.endswith("rfc") or canon == "rfc":
                v = str(v).strip().upper().replace(" ", "") if v else None
                if v and not RFC_RE.match(v):
                    avisos.append(f"{ctx}: '{v}' no tiene forma de RFC")
            elif canon.endswith("clabe"):
                v = re.sub(r"\D", "", str(v)) if v else None
            r[canon] = v

        if not r.get(pk):
            if spec["prefijo"]:
                r[pk] = f"{spec['prefijo']}-{i:05d}"
            elif pk == "entry_id":
                r[pk] = i
            else:
                avisos.append(f"{ctx}: sin {pk}, fila descartada")
                continue
        salida.append(r)
    return salida


# ---------------------------------------------------------------------------
# Rellenos (solo lo deducible; nada inventado)
# ---------------------------------------------------------------------------

def completar_montos(invoices: list[dict], rellenos: list[str],
                     avisos: list[str]) -> list[dict]:
    n_iva = n_sub = n_tot = 0
    # Una factura sin subtotal ni total no se puede reconciliar en pesos, que es
    # justo lo que el validador exige de cada hallazgo. Se descarta con aviso en
    # vez de arrastrar un NULL hasta la mitad del pipeline.
    sin_monto = [r for r in invoices if r.get("subtotal") is None and r.get("total") is None]
    if sin_monto:
        avisos.append(f"{len(sin_monto)} factura(s) sin subtotal ni total: descartadas "
                      f"(sin monto no hay reconciliacion posible). "
                      f"Ej.: {sin_monto[0].get('uuid')}")
        invoices = [r for r in invoices if r not in sin_monto]
    for r in invoices:
        sub, iva, tot = r.get("subtotal"), r.get("iva"), r.get("total")
        if sub is not None and iva is None:
            r["iva"] = round(sub * IVA_RATE, 2); n_iva += 1
        if sub is not None and r.get("total") is None:
            r["total"] = round(sub + (r.get("iva") or 0), 2); n_tot += 1
        elif sub is None and tot is not None:
            r["subtotal"] = round(tot / (1 + IVA_RATE), 2)
            if r.get("iva") is None:
                r["iva"] = round(tot - r["subtotal"], 2)
            n_sub += 1
        if not r.get("status"):
            r["status"] = "vigente"
    if n_iva:
        rellenos.append(f"IVA calculado al {IVA_RATE:.0%} en {n_iva} factura(s)")
    if n_tot:
        rellenos.append(f"total = subtotal + IVA en {n_tot} factura(s)")
    if n_sub:
        rellenos.append(f"subtotal despejado del total en {n_sub} factura(s)")
    return invoices


def derivar_vendors(tablas: dict, rellenos: list[str], avisos: list[str]) -> None:
    """Un proveedor que emitio una factura existe: eso no es inventar. Los
    campos que no conocemos (CLABE, alta, categoria) se dejan en NULL, que
    es honesto; inventarlos haria que un detector cite datos falsos."""
    invoices = tablas.get("invoices", [])
    if not invoices:
        return
    conocidos = {v["rfc"] for v in tablas.get("vendors", []) if v.get("rfc")}
    receptores = [r["receiver_rfc"] for r in invoices if r.get("receiver_rfc")]
    empresa = max(set(receptores), key=receptores.count) if receptores else None

    nuevos = []
    for r in invoices:
        e = r.get("issuer_rfc")
        if e and e != empresa and e not in conocidos:
            conocidos.add(e)
            nuevos.append({"rfc": e, "legal_name": None, "registered_date": None,
                           "address": None, "bank_clabe": None, "category": None,
                           "contact_email": None})
    if nuevos:
        tablas.setdefault("vendors", []).extend(nuevos)
        rellenos.append(f"{len(nuevos)} proveedor(es) deducido(s) de quien emitio "
                        f"cada factura (sin CLABE ni categoria: no se inventan)")
    sin_clabe = sum(1 for v in tablas.get("vendors", []) if not v.get("bank_clabe"))
    if sin_clabe:
        avisos.append(f"{sin_clabe} proveedor(es) sin CLABE: los detectores que "
                      f"cruzan pagos con facturas no podran ligarlos")


def cruzar_69b(tablas: dict, rellenos: list[str], avisos: list[str]) -> None:
    """Cruza los RFC de tu archivo contra el listado 69-B REAL del SAT."""
    if tablas.get("efos_list"):
        return
    if not SAT_CSV.exists():
        avisos.append(f"no encontre {SAT_CSV.relative_to(REPO)}: sin cruce 69-B")
        return
    sys.path.insert(0, str(REPO / "estate_gen"))
    from generate_estate import load_sat_69b  # noqa: E402

    presentes = {v["rfc"] for v in tablas.get("vendors", []) if v.get("rfc")}
    for r in tablas.get("invoices", []):
        for k in ("issuer_rfc", "receiver_rfc"):
            if r.get(k):
                presentes.add(r[k])
    if not presentes:
        return

    # El esquema de efos_list solo contempla definitivo|presunto. Los otros
    # dos estatus reales ('desvirtuado', 'sentencia favorable') son
    # exoneraciones: no pertenecen a una lista de EFOS.
    # load_sat_69b devuelve `status` con el texto crudo del SAT ('Definitivo',
    # 'Sentencia Favorable'...) y `status_norm` ya normalizado. El esquema de
    # efos_list quiere el normalizado.
    hits = [r for r in load_sat_69b(SAT_CSV)
            if r["rfc"] in presentes and r["status_norm"] in ("definitivo", "presunto")]
    if hits:
        tablas["efos_list"] = [{"rfc": r["rfc"], "legal_name": r["legal_name"],
                                "status": r["status_norm"], "publication_date": None}
                               for r in hits]
        defin = sum(1 for r in hits if r["status_norm"] == "definitivo")
        rellenos.append(f"cruce contra el 69-B real del SAT: {len(hits)} RFC de tu "
                        f"archivo estan listados ({defin} definitivo(s))")
    else:
        rellenos.append("cruce contra el 69-B real del SAT: ningun RFC tuyo aparece")


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

def escribir_estate(tablas: dict, salida: Path) -> None:
    salida.parent.mkdir(parents=True, exist_ok=True)
    if salida.exists():
        salida.unlink()
    con = sqlite3.connect(salida)
    con.executescript(SCHEMA_SQL.read_text(encoding="utf-8"))
    for tabla, filas in tablas.items():
        if not filas:
            continue
        cols = TABLAS[tabla]["cols"]
        con.executemany(
            f"INSERT OR REPLACE INTO {tabla} ({','.join(cols)}) "
            f"VALUES ({','.join('?' * len(cols))})",
            [tuple(r.get(c) for c in cols) for r in filas])
    con.commit()
    con.close()


def escribir_plantilla(destino: Path) -> None:
    ejemplos = {
        "vendors": ["AAA010101AA1", "Proveedor Ejemplo SA de CV", "2024-03-15",
                    "Av. Ejemplo 100, Monterrey", "012345678901234567",
                    "Consultoria", "contacto@ejemplo.mx"],
        "invoices": ["INV-00001", "AAA010101AA1", "EMP920101AB1", "2026-02-15",
                     80000.00, 12800.00, 92800.00, "Servicios de consultoria",
                     "G03", "03", "PUE", "vigente"],
        "ledger": [1, "2026-02-15", "5000", "Gastos operativos", 92800.00, 0.00,
                   "Registro factura", "INV-00001", "CC-100", "A. Ejemplo"],
        "bank_txns": ["BNK-00001", "2026-03-01", "099999999999999999",
                      "012345678901234567", 92800.00, "Pago INV-00001", "SPEI"],
        "purchase_orders": ["PO-00001", "AAA010101AA1", "2026-02-10", 92800.00,
                            "C. Ejemplo", "D. Ejemplo", "Consultoria Q1"],
        "contracts": ["CTR-00001", "AAA010101AA1", "2025-01-01", 556800.00,
                      "Contrato marco anual"],
        "employees": ["EMP:0001", "Persona Ejemplo", "Gerente de Compras",
                      "000000000000000501", "2021-03-01"],
        "efos_list": ["AAA010101AA1", "Proveedor Ejemplo SA de CV", "definitivo",
                      "2025-12-11"],
    }
    if destino.suffix.lower() == ".xlsx":
        try:
            from openpyxl import Workbook
        except ImportError:
            raise SystemExit("Para la plantilla .xlsx hace falta openpyxl "
                             "(pip install openpyxl), o pide una carpeta de CSVs: "
                             "--plantilla plantilla_csv/")
        wb = Workbook()
        wb.remove(wb.active)
        for tabla, spec in TABLAS.items():
            ws = wb.create_sheet(tabla)
            ws.append(spec["cols"])
            ws.append(ejemplos[tabla])
            for i, c in enumerate(spec["cols"], 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(c) + 3)
        wb.save(destino)
    else:
        destino.mkdir(parents=True, exist_ok=True)
        for tabla, spec in TABLAS.items():
            with open(destino / f"{tabla}.csv", "w", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                w.writerow(spec["cols"])
                w.writerow(ejemplos[tabla])
    print(f"\nPlantilla escrita en {destino}")
    print("  - Una hoja/archivo por tabla. La fila 2 es un ejemplo: borrala.")
    print("  - Lo minimo util es la hoja 'invoices'. Con 'bank_txns' ademas se "
          "activan los detectores de ciclos y kickbacks.")
    print("  - Puedes renombrar los encabezados al espanol ('RFC Emisor', "
          "'Fecha', 'Monto'): se reconocen igual.")


# ---------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--entrada", type=Path, help=".xlsx, .csv o carpeta con CSVs")
    ap.add_argument("--salida", type=Path, default=REPO / "data" / "estates" / "mi_estate.db")
    ap.add_argument("--tabla", choices=list(TABLAS),
                    help="si la entrada es una sola hoja/CSV, a que tabla corresponde")
    ap.add_argument("--plantilla", type=Path,
                    help="escribe una plantilla vacia (.xlsx o carpeta) y termina")
    ap.add_argument("--sin-69b", action="store_true",
                    help="no cruzar los RFC contra el listado 69-B del SAT")
    ap.add_argument("--puntuar", action="store_true",
                    help="al terminar, corre las etapas deterministas + CART y "
                         "muestra el ranking de sospecha")
    args = ap.parse_args()

    if args.plantilla:
        escribir_plantilla(args.plantilla)
        return 0
    if not args.entrada:
        ap.error("hace falta --entrada (o --plantilla para empezar)")
    if not args.entrada.exists():
        raise SystemExit(f"No existe: {args.entrada}")

    avisos: list[str] = []
    rellenos: list[str] = []

    hojas = cargar_entrada(args.entrada, args.tabla)
    tablas: dict[str, list[dict]] = {}
    print(f"\nLeyendo {args.entrada.name}")
    for nombre, filas in hojas.items():
        tabla = args.tabla if args.tabla and len(hojas) == 1 else emparejar_hoja(nombre)
        if not tabla:
            avisos.append(f"hoja '{nombre}': no corresponde a ninguna tabla del "
                          f"esquema, se ignoro")
            continue
        convertidas = convertir_tabla(tabla, filas, avisos)
        tablas.setdefault(tabla, []).extend(convertidas)
        print(f"  hoja '{nombre}' -> {tabla}: {len(convertidas)} fila(s)")

    if not tablas.get("invoices"):
        raise SystemExit(
            "\nNo se leyo ninguna factura. El pipeline construye el perfil de cada "
            "entidad a partir de sus facturas; sin esa tabla no hay nada que "
            "puntuar.\nSaca la plantilla con --plantilla plantilla.xlsx para ver "
            "las columnas esperadas.")

    tablas["invoices"] = completar_montos(tablas["invoices"], rellenos, avisos)
    derivar_vendors(tablas, rellenos, avisos)
    if not args.sin_69b:
        cruzar_69b(tablas, rellenos, avisos)

    escribir_estate(tablas, args.salida)

    print(f"\nEstate escrita: {args.salida}")
    for tabla in TABLAS:
        n = len(tablas.get(tabla, []))
        print(f"  {tabla:18s} {n:>6}" + ("" if n else "   (vacia)"))

    if rellenos:
        print("\nRellenado (deducido, no inventado):")
        for r in rellenos:
            print(f"  - {r}")

    faltantes = [t for t in ("bank_txns", "purchase_orders", "ledger", "employees")
                 if not tablas.get(t)]
    if faltantes:
        print("\nOjo — tablas vacias y lo que dejan de detectar:")
        detalle = {
            "bank_txns": "ciclos de transferencias (round_tripping), kickbacks y "
                         "pagos no rastreables",
            "purchase_orders": "fraccionamiento bajo el umbral de 50k y "
                               "solicitante = aprobador",
            "ledger": "facturas sin respaldo contable",
            "employees": "cualquier esquema que ligue a un empleado",
        }
        for t in faltantes:
            print(f"  - sin {t}: no se evalua {detalle[t]}")
        print("  Esto no es un error: el sistema solo acusa con evidencia que "
              "exista en la base.")

    if avisos:
        print(f"\nAvisos ({len(avisos)}):")
        for a in avisos[:25]:
            print(f"  - {a}")
        if len(avisos) > 25:
            print(f"  ... y {len(avisos) - 25} mas")

    if args.puntuar:
        from src.scoring import generar_leads
        from src.tools import EstateDB
        with EstateDB(args.salida) as estate:
            empresa = estate.identificar_empresa()
            print(f"\nEmpresa auditada (inferida): RFC:{empresa.rfc} "
                  f"(confianza {empresa.confianza_rfc:.0%})")
            leads = generar_leads(estate)
            print(f"\n{len(leads)} entidad(es) con senales, de mayor a menor sospecha:\n")
            print(f"  {'#':>3} {'entidad':24s} {'score':>9s}  {'modelo (CART)':20s} senales")
            print("  " + "-" * 92)
            for i, l in enumerate(leads[:15], 1):
                dets = ", ".join(sorted({s.detector.replace("detectar_", "")
                                         for s in l.signals}))
                print(f"  {i:>3} {l.entity:24s} {l.score:9.6f}  "
                      f"{str(l.ml_scheme_type):20s} {dets[:44]}")
            if not leads:
                print("  ninguna. Con los datos dados, nada disparo una regla.")

    print(f"\nSiguiente paso — pipeline completo con Gemma:\n"
          f"  python3 -m src.run --estate {args.salida} --out salida.json")
    print(f"Ver por que quedo cada entidad donde quedo:\n"
          f"  python3 scripts/explicar.py --estate {args.salida} --entidad RFC:XXXX")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
