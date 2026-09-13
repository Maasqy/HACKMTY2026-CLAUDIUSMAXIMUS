#!/usr/bin/env python3
"""
scripts/generar_ejemplo_proveedores.py — un Excel de PRUEBA para
scripts/excel_a_estate.py, sin depender de tener datos reales a la mano.

Simula lo que alguien de finanzas exportaria de su ERP: facturas, pagos,
ordenes de compra, contratos y nomina de una empresa con seis proveedores.
Cuatro son de rutina. Dos traen un patron sembrado a proposito para que el
pipeline tenga algo que encontrar al probarlo:

  - un PHANTOM_VENDOR: un RFC REAL del listado 69-B del SAT (status
    'definitivo'), facturando puros conceptos genericos, sin orden de
    compra ni contrato que lo respalde.
  - un KICKBACK: un proveedor normal cuyo pago la empresa deposita, y que
    2-3 dias despues retransfiere ~85% a la cuenta del mismo empleado que
    aprueba (y solicita) todas sus ordenes de compra.

Uso:
    python3 scripts/generar_ejemplo_proveedores.py
    python3 scripts/generar_ejemplo_proveedores.py --salida otro_nombre.xlsx

Con el archivo en mano:
    python3 scripts/excel_a_estate.py --entrada mis_proveedores.xlsx \\
        --salida data/estates/mis_proveedores.db --puntuar

    python3 -m src.run --estate data/estates/mis_proveedores.db \\
        --out salida.json --max-leads 4

    python3 scripts/reporte_excel.py --submission salida.json \\
        --estate data/estates/mis_proveedores.db --empresa "Servicios Preferentes"
"""

from __future__ import annotations

import argparse
from datetime import date, timedelta
from pathlib import Path

COMPANY_CLABE = "012180001234567890"

PROVEEDORES = [
    # RFC, Razón Social, CLABE, Giro, Fecha de Alta
    ("CFL180312AB3", "Consultoria Fiscal del Levante SA de CV", "014180009876543210", "Consultoria", "2022-04-10"),
    ("MRN150920TU8", "Mantenimiento Regional del Norte SC", "021180003344556677", "Mantenimiento", "2021-08-01"),
    ("TLC090114QW2", "Transportes y Logistica del Centro SA de CV", "072180001122334455", "Transporte", "2020-02-15"),
    ("PAP110603RE5", "Papeleria y Suministros PROOFICE SA de CV", "044180005566778899", "Papeleria", "2019-11-20"),
    # Fantasma: RFC REAL del 69-B ("definitivo").
    ("AAA120730823", "Asesores y Administradores Agricolas SA", "058180001231231234", "Consultoria", "2026-01-05"),
    # Kickback: factura normal, pero el pago se retransfiere al aprobador.
    ("SVP200815KL9", "Servicios Preferentes de Ponente SA de CV", "127180009998887776", "Consultoria", "2025-09-12"),
]
CLABE_EMPLEADO = "127180001112223334"


def construir(salida: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Proveedores")
    ws.append(["RFC", "Razon Social", "CLABE", "Giro", "Fecha de Alta"])
    for p in PROVEEDORES:
        ws.append(list(p))

    ws = wb.create_sheet("Empleados")
    ws.append(["ID", "Nombre", "Puesto", "CLABE", "Fecha de Ingreso"])
    ws.append(["EMP:0007", "Ricardo Elizondo Sada", "Gerente de Compras", CLABE_EMPLEADO, "2019-03-01"])

    clabes = {p[0]: p[2] for p in PROVEEDORES}
    base = date(2026, 1, 15)
    facturas: list[tuple[str, str, date, float]] = []
    folio = 1

    ws_f = wb.create_sheet("Facturas")
    ws_f.append(["Folio", "RFC Emisor", "RFC Receptor", "Fecha de Emision",
                "Subtotal", "Concepto", "Metodo de Pago", "Estatus"])

    def add_inv(rfc, fecha, subtotal, concepto):
        nonlocal folio
        fid = f"F-{folio:04d}"
        ws_f.append([fid, rfc, "EMP920101AB1", fecha.isoformat(), subtotal, concepto, "PUE", "vigente"])
        facturas.append((fid, rfc, fecha, subtotal))
        folio += 1

    add_inv("CFL180312AB3", base, 42000.00, "Dictamen fiscal ejercicio 2025")
    add_inv("CFL180312AB3", base + timedelta(days=45), 38500.00, "Revision de nomina Q1 2026")
    add_inv("MRN150920TU8", base + timedelta(days=5), 61200.00, "Mantenimiento preventivo planta Monterrey")
    add_inv("MRN150920TU8", base + timedelta(days=70), 58900.00, "Reparacion compresor linea 2")
    add_inv("TLC090114QW2", base + timedelta(days=10), 27300.00, "Flete Monterrey-Saltillo, 12 viajes")
    add_inv("TLC090114QW2", base + timedelta(days=40), 31150.00, "Flete Monterrey-Reynosa, 14 viajes")
    add_inv("PAP110603RE5", base + timedelta(days=20), 8400.00, "Consumibles de oficina marzo")
    add_inv("PAP110603RE5", base + timedelta(days=80), 9100.00, "Consumibles de oficina mayo")
    add_inv("AAA120730823", base + timedelta(days=8), 96000.00, "Servicios profesionales diversos")
    add_inv("AAA120730823", base + timedelta(days=38), 87500.00, "Servicios de consultoria general")
    add_inv("AAA120730823", base + timedelta(days=66), 91200.00, "Honorarios por servicios prestados")
    add_inv("SVP200815KL9", base + timedelta(days=12), 118000.00, "Asesoria estrategica primer trimestre")
    add_inv("SVP200815KL9", base + timedelta(days=52), 97400.00, "Asesoria estrategica segundo trimestre")
    add_inv("SVP200815KL9", base + timedelta(days=95), 84900.00, "Asesoria estrategica tercer trimestre")

    ws_t = wb.create_sheet("Transferencias")
    ws_t.append(["Folio", "Fecha", "CLABE Origen", "CLABE Destino", "Monto", "Referencia", "Canal"])
    txn = 1

    def add_txn(fecha, origen, destino, monto, ref):
        nonlocal txn
        ws_t.append([f"T-{txn:04d}", fecha.isoformat(), origen, destino, monto, ref, "SPEI"])
        txn += 1

    for fid, rfc, fecha, subtotal in facturas:
        total = round(subtotal * 1.16, 2)
        add_txn(fecha + timedelta(days=10), COMPANY_CLABE, clabes[rfc], total, f"Pago {fid}")

    for fid, rfc, fecha, subtotal in [f for f in facturas if f[1] == "SVP200815KL9"]:
        total = round(subtotal * 1.16, 2)
        add_txn(fecha + timedelta(days=13), clabes["SVP200815KL9"], CLABE_EMPLEADO,
                round(total * 0.85, 2), "Reembolso de gastos")

    ws_po = wb.create_sheet("Ordenes de Compra")
    ws_po.append(["Orden", "RFC Proveedor", "Fecha", "Monto", "Solicitante", "Aprobador", "Descripcion"])
    po = 1
    for fid, rfc, fecha, subtotal in [f for f in facturas if f[1] == "SVP200815KL9"]:
        ws_po.append([f"PO-{po:04d}", rfc, fecha.isoformat(), round(subtotal * 1.16, 2),
                     "Ricardo Elizondo Sada", "Ricardo Elizondo Sada", "Asesoria estrategica trimestral"])
        po += 1
    solicitantes = [("Laura Benavides", "Jorge Salinas"), ("Marco Villalobos", "Jorge Salinas")]
    for rfc in ("CFL180312AB3", "MRN150920TU8", "TLC090114QW2"):
        for i, (fid, r2, fecha, subtotal) in enumerate([f for f in facturas if f[1] == rfc]):
            s, a = solicitantes[i % len(solicitantes)]
            ws_po.append([f"PO-{po:04d}", rfc, fecha.isoformat(), round(subtotal * 1.16, 2), s, a,
                         "Servicio recurrente"])
            po += 1

    ws_c = wb.create_sheet("Contratos")
    ws_c.append(["Contrato", "RFC Proveedor", "Fecha de Inicio", "Valor", "Alcance"])
    ws_c.append(["CTR-0001", "CFL180312AB3", "2022-04-10", 480000.00, "Servicios fiscales anuales"])
    ws_c.append(["CTR-0002", "MRN150920TU8", "2021-08-01", 720000.00, "Mantenimiento industrial anual"])
    ws_c.append(["CTR-0003", "TLC090114QW2", "2020-02-15", 372000.00, "Fletes foraneos anuales"])

    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    print(f"Escrito: {salida}")
    print(f"  {len(PROVEEDORES)} proveedores, {len(facturas)} facturas, "
          f"{txn - 1} transferencias, {po - 1} ordenes")
    print("\nCasos sembrados a proposito:")
    print("  RFC:AAA120730823  phantom_vendor  (69-B real, sin PO ni contrato)")
    print("  RFC:SVP200815KL9  kickback        (retransfiere al aprobador de sus PO)")
    print(f"\nSiguiente paso:\n"
          f"  python3 scripts/excel_a_estate.py --entrada {salida} "
          f"--salida data/estates/mis_proveedores.db --puntuar")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--salida", type=Path, default=Path("mis_proveedores.xlsx"))
    args = ap.parse_args()
    try:
        construir(args.salida)
    except ImportError:
        raise SystemExit("Hace falta openpyxl:  pip install openpyxl")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
